#!/usr/bin/env python3
"""Build the 10-slide PowerPoint for the RKB Restaurant Order System."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(BASE, "images")
OUT = "/home/hatch/workspace/your_files/restaurant-order-presentation.pptx"

NAVY = RGBColor(0x0A, 0x28, 0x5A)
BLUE = RGBColor(0x00, 0x77, 0xB5)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
DARK = RGBColor(0x1E, 0x1E, 0x1E)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, runs, size=24, bold=False,
             color=DARK, alignment=PP_ALIGN.LEFT):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    tf = tx.text_frame
    tf.word_wrap = True
    for i, text in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alignment
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tx


def title_slide(slide, title, subtitle):
    bg(slide, NAVY)
    add_text(slide, 0.8, 1.6, 11.7, 2, [title], size=54, bold=True,
             color=WHITE, alignment=PP_ALIGN.CENTER)
    add_text(slide, 0.8, 3.8, 11.7, 1.5, [subtitle], size=28,
             color=RGBColor(0xBB, 0xDD, 0xFF), alignment=PP_ALIGN.CENTER)


def content_slide(title, bullets, note=None):
    slide = prs.slides.add_slide(blank)
    bg(slide)
    add_text(slide, 0.7, 0.25, 11.9, 1, [title], size=40, bold=True, color=BLUE)
    tx = slide.shapes.add_textbox(Inches(0.9), Inches(1.5),
                                  Inches(11.5), Inches(5))
    tf = tx.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        p.level = 0
        r = p.add_run()
        r.text = "•  " + b
        r.font.size = Pt(24)
        r.font.color.rgb = DARK
    if note:
        add_text(slide, 0.9, 6.5, 11.5, 0.7, [note], size=18, color=GRAY)
    return slide


def image_slide(title, img_name, caption=None):
    slide = prs.slides.add_slide(blank)
    bg(slide)
    add_text(slide, 0.7, 0.25, 11.9, 1, [title], size=40, bold=True, color=BLUE)
    slide.shapes.add_picture(os.path.join(IMG, img_name),
                             Inches(0.6), Inches(1.35),
                             width=Inches(12.1), height=Inches(5.35))
    if caption:
        add_text(slide, 0.7, 6.85, 11.9, 0.5, [caption], size=18, color=GRAY)
    return slide


# 1 — title
s = prs.slides.add_slide(blank)
title_slide(s, "RKB — Restaurant Order\nManagement System",
            "Python · FastAPI · Raw SQL (no ORM)\nOrder lifecycle, kitchen flow & polite delay handling")

# 2 — what it does
content_slide("What it does", [
    "Full order lifecycle for a restaurant: waiter takes the order, kitchen "
    "prepares it, waiter serves it, customer care handles delays",
    "Strict order state machine: PLACED → KITCHEN_RECEIVED → PREPARING → "
    "READY → SERVED → COMPLETED (plus CANCELLED)",
    "Invalid transitions (skipping or reversing states) are rejected with HTTP 400",
    "Written SQL-first: every query is explicit raw SQL — no ORM",
])

# 3 — order flow diagram
image_slide("Order flow", "order_flow.png",
            "Invalid moves (e.g. PLACED → READY) are rejected; CANCELLED is reachable from any non-final state.")

# 4 — roles
content_slide("Three roles, one flow", [
    "WAITER — takes orders (table, customer name/phone, items), serves ready "
    "orders to tables, completes them after payment",
    "KITCHEN — watches the live kitchen queue (oldest first), starts preparing, "
    "marks orders ready",
    "CUSTOMER CARE — watches delayed orders, sends a polite delay notice, "
    "records the customer's response",
    "Placing an order occupies the table; completing or cancelling frees it",
])

# 5 — delay handling (key feature)
content_slide("Key feature: polite delay handling", [
    "Each menu item has an expected_prep_minutes value",
    "An order in PREPARING becomes DELAYED when it exceeds the max expected "
    "prep time of its items (floor: 15 minutes, configurable)",
    "Customer care sends a courteous templated message:",
    "\u201cWe're sorry for the wait! Your order will take 10 to 15 more minutes. "
    "Thank you for your patience.\u201d",
    "The customer's reply is recorded: WAITING, CANCELLED_ORDER, ESCALATED, "
    "or NO_RESPONSE",
])

# 6 — tech stack
content_slide("Tech stack", [
    "API: FastAPI + Pydantic v2",
    "Database: raw SQL via app/db/db.py — zero ORM",
    "Drivers: sqlite3 (stdlib) for dev & tests, psycopg2 for PostgreSQL in prod",
    "One SQL dialect: ? placeholders auto-translated to %s for psycopg2; "
    "timestamps stored as UTC ISO-8601 text on both engines",
    "Tests: pytest + FastAPI TestClient (27 tests, in-memory SQLite)",
    "Deploy: Docker + Docker Compose (PostgreSQL 16)",
])

# 7 — architecture diagram
image_slide("Architecture", "architecture.png",
            "Requests flow through routers → services (state machine, delay detector, notifier) → raw SQL.")

# 8 — data model diagram
image_slide("Data model", "data_model.png",
            "Five tables; order_items snapshots price & prep time at order time so history never changes.")

# 9 — API + tests
content_slide("API & quality", [
    "POST /api/orders — take an order;  GET /api/orders?status= — filter",
    "POST /api/orders/{id}/advance {to_status, actor_role} — state machine",
    "GET /api/kitchen/queue — live kitchen work list",
    "GET /api/customer-care/delayed — most overdue first",
    "POST /api/customer-care/{id}/notify — send polite delay message",
    "POST /api/customer-care/{id}/respond — record WAITING / CANCELLED_ORDER / …",
    "GET /api/dashboard/stats — counts, avg prep minutes, delayed, revenue",
    "27 automated tests, all passing: happy path, invalid transitions, delay "
    "detection, notify/respond, CRUD, stats",
])

# 10 — closing
s = prs.slides.add_slide(blank)
bg(s, NAVY)
add_text(s, 0.8, 1.4, 11.7, 2, ["Try it yourself"], size=54, bold=True,
         color=WHITE, alignment=PP_ALIGN.CENTER)
add_text(s, 0.8, 3.4, 11.7, 3,
         ["Code: github.com/rpoluri9/restaurant-order-system",
          "",
          "python seed.py && uvicorn app.main:app --reload",
          "Open http://127.0.0.1:8000/docs",
          "",
          "Thank you!"], size=26,
         color=RGBColor(0xBB, 0xDD, 0xFF), alignment=PP_ALIGN.CENTER)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved", OUT)
